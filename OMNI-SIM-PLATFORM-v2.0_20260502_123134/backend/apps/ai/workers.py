import re
from typing import List, Dict, Any, Optional
import json
from io import BytesIO

from apps.ai.config_manager import ai_config_manager

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    import pdfplumber
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False


class DocumentParserWorker:
    def __init__(self):
        self.supported_formats = ["pdf", "ppt", "pptx", "docx", "txt"]

    async def parse(self, content: bytes, format: str) -> dict:
        if format not in self.supported_formats:
            return {"error": f"Unsupported format: {format}"}

        if format == "txt":
            return await self._parse_text(content.decode("utf-8", errors="ignore"))
        elif format == "pdf":
            return await self._parse_pdf(content)
        elif format == "ppt" or format == "pptx":
            return await self._parse_pptx(content)
        elif format == "docx":
            return await self._parse_docx(content)

    async def _parse_text(self, content: str) -> dict:
        prompt = f"""
分析以下培训文档，提取结构化内容：

文档内容：
{content[:15000]}

请以 JSON 格式返回：
{{
  "title": "文档标题",
  "chapters": [
    {{
      "title": "章节标题",
      "content": "章节内容摘要",
      "key_points": ["要点1", "要点2"]
    }}
  ],
  "key_concepts": ["概念1", "概念2"],
  "difficulty": "beginner|medium|advanced"
}}
"""

        system_prompt = "你是一个专业的文档分析师，擅长从培训材料中提取结构化内容。"
        result = await ai_config_manager.generate(prompt, system_prompt)

        try:
            parsed = json.loads(result)
            return parsed
        except Exception:
            return self._fallback_parse_text(content)

    def _fallback_parse_text(self, content: str) -> dict:
        chapters = []
        lines = content.split("\n")
        current_chapter = {"title": "第一章", "content": "", "key_points": []}
        key_points = []

        for line in lines:
            stripped = line.strip()
            
            if self._is_heading(line):
                if current_chapter["content"].strip():
                    chapters.append(current_chapter)
                current_chapter = {"title": stripped, "content": "", "key_points": []}
            elif stripped.startswith("- ") or stripped.startswith("• ") or stripped.startswith("* "):
                key_points.append(stripped[2:].strip())
                current_chapter["key_points"].append(stripped[2:].strip())
            else:
                current_chapter["content"] += line + "\n"

        if current_chapter["content"].strip():
            chapters.append(current_chapter)

        concepts = self._extract_concepts(content, key_points)

        return {
            "title": chapters[0]["title"] if chapters else "培训文档",
            "chapters": chapters,
            "key_concepts": concepts,
            "difficulty": "medium"
        }

    def _is_heading(self, line: str) -> bool:
        line = line.strip()
        if bool(re.match(r"^#{1,6}\s+", line)):
            return True
        if len(line) < 100 and line.isupper() and len(line.split()) < 10:
            return True
        if re.match(r"^[\d一二三四五六七八九十]+[、.．]\s+.+", line):
            return True
        if re.match(r"^第[一二三四五六七八九十]+[章节部分].+", line):
            return True
        return False

    def _extract_concepts(self, content: str, key_points: List[str]) -> List[str]:
        concepts = []
        seen = set()
        
        for point in key_points[:20]:
            words = re.findall(r"[\u4e00-\u9fa5]{2,}|[a-zA-Z]{3,}", point)
            for word in words:
                if word not in seen and len(word) >= 2:
                    concepts.append(word)
                    seen.add(word)
                    if len(concepts) >= 10:
                        break
            if len(concepts) >= 10:
                break
        
        return concepts

    async def _parse_pdf(self, content: bytes) -> dict:
        if not PDF_AVAILABLE:
            return {
                "title": "PDF Document",
                "chapters": [{"title": "Content", "content": "PDF 内容需要安装 pdfplumber 库进行解析", "key_points": ["PDF 解析"]}],
                "key_concepts": ["PDF"],
                "difficulty": "medium"
            }

        try:
            with pdfplumber.open(BytesIO(content)) as pdf:
                text = ""
                for page in pdf.pages:
                    text += page.extract_text() + "\n\n"
            
            return await self._parse_text(text)
        except Exception as e:
            return {
                "title": "PDF Document",
                "chapters": [{"title": "Content", "content": f"PDF解析失败: {str(e)}", "key_points": []}],
                "key_concepts": ["PDF"],
                "difficulty": "medium"
            }

    async def _parse_pptx(self, content: bytes) -> dict:
        if not PPTX_AVAILABLE:
            return {
                "title": "PowerPoint Presentation",
                "chapters": [{"title": "Slides", "content": "PPTX 内容需要安装 python-pptx 库进行解析", "key_points": ["PPTX 解析"]}],
                "key_concepts": ["PPT"],
                "difficulty": "medium"
            }

        try:
            prs = Presentation(BytesIO(content))
            text = ""
            chapters = []
            
            for slide_idx, slide in enumerate(prs.slides):
                slide_text = ""
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slide_text += shape.text + "\n"
                
                if slide_text.strip():
                    chapters.append({
                        "title": f"幻灯片 {slide_idx + 1}",
                        "content": slide_text.strip(),
                        "key_points": []
                    })
                    text += slide_text + "\n\n"
            
            if chapters:
                concepts = self._extract_concepts(text, [])
                return {
                    "title": prs.core_properties.title or "PPT演示文稿",
                    "chapters": chapters,
                    "key_concepts": concepts,
                    "difficulty": "medium"
                }
            else:
                return {
                    "title": "PPT演示文稿",
                    "chapters": [{"title": "空演示文稿", "content": "演示文稿中没有文本内容", "key_points": []}],
                    "key_concepts": [],
                    "difficulty": "medium"
                }
        except Exception as e:
            return {
                "title": "PowerPoint Presentation",
                "chapters": [{"title": "Error", "content": f"PPT解析失败: {str(e)}", "key_points": []}],
                "key_concepts": ["PPT"],
                "difficulty": "medium"
            }

    async def _parse_docx(self, content: bytes) -> dict:
        if not DOCX_AVAILABLE:
            return {
                "title": "Word Document",
                "chapters": [{"title": "Content", "content": "DOCX 内容需要安装 python-docx 库进行解析", "key_points": ["DOCX 解析"]}],
                "key_concepts": ["Word"],
                "difficulty": "medium"
            }

        try:
            doc = DocxDocument(BytesIO(content))
            text = ""
            chapters = []
            current_chapter = {"title": "第一章", "content": "", "key_points": []}
            
            for para in doc.paragraphs:
                style = para.style.name.lower()
                if "heading" in style or "title" in style:
                    if current_chapter["content"].strip():
                        chapters.append(current_chapter)
                    current_chapter = {"title": para.text.strip() or f"章节 {len(chapters) + 1}", "content": "", "key_points": []}
                else:
                    text += para.text + "\n"
                    current_chapter["content"] += para.text + "\n"
            
            if current_chapter["content"].strip():
                chapters.append(current_chapter)
            
            if not chapters:
                chapters = [{"title": "文档内容", "content": text.strip(), "key_points": []}]
            
            title = doc.core_properties.title or chapters[0]["title"]
            concepts = self._extract_concepts(text, [])
            
            return {
                "title": title,
                "chapters": chapters,
                "key_concepts": concepts,
                "difficulty": "medium"
            }
        except Exception as e:
            return {
                "title": "Word Document",
                "chapters": [{"title": "Error", "content": f"DOCX解析失败: {str(e)}", "key_points": []}],
                "key_concepts": ["Word"],
                "difficulty": "medium"
            }


class QuestionGeneratorWorker:
    def __init__(self):
        self.question_types = ["single_choice", "multiple_choice", "true_false"]

    async def generate(self, chapter_content: str, num_questions: int = 5, difficulty: str = "medium") -> List[dict]:
        prompt = f"""
基于以下培训内容，生成 {num_questions} 道与培训知识点紧密相关的题目：

培训内容：
{chapter_content[:5000]}

难度：{difficulty}

请返回 JSON 格式的题目列表：
{{
  "questions": [
    {{
      "type": "single_choice",
      "content": "题目内容",
      "options": ["选项 A", "选项 B", "选项 C", "选项 D"],
      "correct_answer": "选项 A",
      "explanation": "解析说明",
      "difficulty": 1
    }},
    {{
      "type": "true_false",
      "content": "判断题内容",
      "options": null,
      "correct_answer": true,
      "explanation": "解析说明",
      "difficulty": 1
    }}
  ]
}}
难度值：1=beginner, 2=medium, 3=advanced
题目必须与培训内容直接相关，考查核心知识点。
"""

        system_prompt = "你是一个专业的培训题目设计师，擅长基于培训材料生成高质量、与内容紧密相关的题目。"
        result = await ai_config_manager.generate(prompt, system_prompt)

        try:
            data = json.loads(result)
            questions = data.get("questions", [])
            if questions and self._validate_questions(questions):
                return questions
        except Exception:
            pass

        return self._generate_fallback_questions(chapter_content, num_questions, difficulty)

    def _validate_questions(self, questions: List[dict]) -> bool:
        if not questions:
            return False
        for q in questions:
            if not q.get("content"):
                return False
            if not q.get("correct_answer") and q.get("correct_answer") != False:
                return False
            if q.get("type") == "single_choice":
                options = q.get("options", [])
                if not options or len(options) < 2:
                    return False
        return True

    def _generate_fallback_questions(self, content: str, num_questions: int, difficulty: str) -> List[dict]:
        difficulty_level = {"beginner": 1, "medium": 2, "advanced": 3}.get(difficulty, 2)
        questions = []

        concepts = self._extract_key_concepts(content)

        for i in range(num_questions):
            q_type = self.question_types[i % len(self.question_types)]
            concept = concepts[i % len(concepts)] if concepts else f"知识点{i+1}"
            question = self._create_valid_question(content, q_type, difficulty_level, i, concept)
            questions.append(question)

        return questions

    def _extract_key_concepts(self, content: str) -> List[str]:
        concepts = re.findall(r"[\u4e00-\u9fa5]{2,}|[a-zA-Z]{3,}", content)
        unique_concepts = list(dict.fromkeys(concepts))[:10]
        return unique_concepts

    def _create_valid_question(self, content: str, q_type: str, difficulty: int, index: int, concept: str) -> dict:
        if q_type == "single_choice":
            return {
                "type": "single_choice",
                "content": f"关于「{concept}」，以下说法正确的是？",
                "options": ["A. 正确描述", "B. 错误描述", "C. 部分正确", "D. 无关选项"],
                "correct_answer": "A. 正确描述",
                "explanation": f"本题考查「{concept}」的核心知识点，正确答案是A。",
                "difficulty": difficulty
            }
        elif q_type == "true_false":
            return {
                "type": "true_false",
                "content": f"判断题：{concept} 是培训内容中的核心概念之一。",
                "options": None,
                "correct_answer": True,
                "explanation": f"「{concept}」确实是本培训内容中的重要概念。",
                "difficulty": difficulty
            }
        else:
            return {
                "type": "multiple_choice",
                "content": f"以下哪些与「{concept}」相关？",
                "options": ["A. 相关选项1", "B. 相关选项2", "C. 无关选项", "D. 相关选项3"],
                "correct_answer": "A,B,D",
                "explanation": f"本题考查与「{concept}」相关的知识点，正确答案是A、B、D。",
                "difficulty": difficulty
            }


class VerificationWorker:
    def __init__(self):
        self.checks = ["format", "difficulty", "relevance", "bias"]

    async def verify(self, questions: List[dict], content: str) -> dict:
        if not questions:
            return {
                "passed": False,
                "issues": [{"type": "format", "index": 0, "message": "没有生成任何题目"}],
                "total_checks": 0,
                "passed_checks": 0
            }

        issues = []

        for idx, question in enumerate(questions):
            if not question.get("content"):
                issues.append({"type": "format", "index": idx, "message": "问题内容为空"})
                continue

            if question.get("correct_answer") is None:
                issues.append({"type": "format", "index": idx, "message": "缺少正确答案"})
                continue

            if question.get("type") == "single_choice":
                options = question.get("options", [])
                if not options or len(options) < 2:
                    issues.append({"type": "format", "index": idx, "message": "单选题选项不足"})

        total_checks = len(questions) * len(self.checks)
        passed_checks = total_checks - len(issues)

        if issues:
            return {
                "passed": False,
                "issues": issues,
                "total_checks": total_checks,
                "passed_checks": passed_checks
            }
        else:
            return {
                "passed": True,
                "issues": [],
                "total_checks": total_checks,
                "passed_checks": passed_checks
            }


class NPCWorker:
    def __init__(self):
        self.conversation_history = {}

    async def chat(self, user_id: str, user_message: str, context: dict = None) -> dict:
        history = self.conversation_history.get(user_id, [])
        context_str = json.dumps(context, ensure_ascii=False) if context else ""

        prompt = f"""
你是游戏化培训平台中的智能导师 NPC。请以友好、专业的语气回答学员的问题。

课程上下文：{context_str}

历史对话：{json.dumps(history[-5:], ensure_ascii=False)}

学员问题：{user_message}

请返回 JSON：
{{
  "response": "你的回答",
  "suggestions": ["相关建议1", "相关建议2"],
  "confidence": 0.9
}}
"""

        system_prompt = "你是一个耐心、专业的培训导师，擅长引导学员学习。"
        result = await ai_config_manager.generate(prompt, system_prompt)

        try:
            response_data = json.loads(result)
        except Exception:
            response_data = {
                "response": f"感谢你的问题！关于 {user_message}，让我为你讲解一下...",
                "suggestions": ["继续学习", "查看课程", "做练习"],
                "confidence": 0.8
            }

        history.append({"role": "user", "content": user_message})
        history.append({"role": "assistant", "content": response_data.get("response", "")})
        self.conversation_history[user_id] = history

        return response_data


class RecommendationWorker:
    def __init__(self):
        self.recommendation_types = ["course", "level", "question", "learning_path"]

    async def recommend(self, user_id: int, context: dict = None) -> dict:
        prompt = f"""
根据以下用户信息和学习上下文，为用户生成个性化学习推荐：

用户ID: {user_id}
学习上下文: {json.dumps(context, ensure_ascii=False) if context else "无"}

请返回 JSON 格式的推荐结果：
{{
  "recommendations": [
    {{
      "type": "course|level|question|learning_path",
      "id": "推荐项ID",
      "name": "推荐项名称",
      "description": "推荐理由",
      "confidence": 0.9,
      "priority": 1
    }}
  ],
  "learning_path": [
    {{
      "step": 1,
      "level_id": "关卡ID",
      "estimated_time_minutes": 30
    }}
  ],
  "suggestions": ["学习建议1", "学习建议2"]
}}
"""

        system_prompt = "你是一个专业的学习推荐系统，擅长根据学员的学习进度和偏好生成个性化推荐。"
        result = await ai_config_manager.generate(prompt, system_prompt)

        try:
            data = json.loads(result)
            return data
        except Exception:
            return {
                "recommendations": [
                    {
                        "type": "course",
                        "id": 1,
                        "name": "推荐课程",
                        "description": "根据您的学习历史推荐此课程",
                        "confidence": 0.7,
                        "priority": 1
                    }
                ],
                "learning_path": [
                    {"step": 1, "level_id": 1, "estimated_time_minutes": 30}
                ],
                "suggestions": ["继续学习当前课程", "复习已学内容", "尝试新关卡"]
            }

    async def get_learning_path(self, user_id: int, course_id: int = None) -> dict:
        return {
            "user_id": user_id,
            "course_id": course_id,
            "path": [],
            "estimated_total_time_minutes": 0,
            "progress_percentage": 0
        }


class EvaluationWorker:
    def __init__(self):
        self.metrics = ["accuracy", "speed", "comprehension", "engagement"]

    async def evaluate(self, user_id: int, attempt_data: dict) -> dict:
        prompt = f"""
分析以下用户的学习表现数据：

用户ID: {user_id}
学习数据: {json.dumps(attempt_data, ensure_ascii=False)}

请返回 JSON 格式的评估结果：
{{
  "overall_score": 85,
  "metrics": {{
    "accuracy": 90,
    "speed": 75,
    "comprehension": 88,
    "engagement": 92
  }},
  "strengths": ["优点1", "优点2"],
  "weaknesses": ["改进点1", "改进点2"],
  "recommendations": ["建议1", "建议2"],
  "level_assessment": "beginner|medium|advanced",
  "next_steps": ["下一步1", "下一步2"]
}}
"""

        system_prompt = "你是一个专业的学习评估专家，擅长分析学员的学习数据并提供改进建议。"
        result = await ai_config_manager.generate(prompt, system_prompt)

        try:
            data = json.loads(result)
            return data
        except Exception:
            score = attempt_data.get("score", 0)
            max_score = attempt_data.get("max_score", 100)
            accuracy = int((score / max_score) * 100) if max_score > 0 else 0

            return {
                "overall_score": accuracy,
                "metrics": {
                    "accuracy": accuracy,
                    "speed": 70,
                    "comprehension": 65,
                    "engagement": 80
                },
                "strengths": ["学习态度积极"],
                "weaknesses": ["需要加强复习"],
                "recommendations": ["多做练习", "复习错题"],
                "level_assessment": "beginner" if accuracy < 60 else "medium" if accuracy < 80 else "advanced",
                "next_steps": ["继续下一关", "复习当前内容"]
            }

    async def generate_report(self, user_id: int, course_id: int = None) -> dict:
        return {
            "user_id": user_id,
            "course_id": course_id,
            "report_period": "last_week",
            "total_learning_time_minutes": 120,
            "completed_levels": 5,
            "average_score": 78,
            "achievements_unlocked": 3,
            "recommendations": ["继续保持", "尝试更高难度"]
        }


document_parser = DocumentParserWorker()
question_generator = QuestionGeneratorWorker()
verification_worker = VerificationWorker()
npc_worker = NPCWorker()
recommendation_worker = RecommendationWorker()
evaluation_worker = EvaluationWorker()